import os
import re
from typing import List, Dict, Any, Tuple
from ingestion import ExtractedContent, NarrativeBlock, TableBlock, VisualBlock, DocumentStats

def table_to_markdown(table: List[List[Any]]) -> str:
    if not table or not table[0]:
        return ""
    clean_table = []
    for row in table:
        clean_table.append([str(cell or "").strip().replace("\n", " ") for cell in row])
        
    headers = clean_table[0]
    if not any(headers):
        headers = [f"Column {i+1}" for i in range(len(headers))]
        
    markdown = f"| {' | '.join(headers)} |\n"
    markdown += f"| {' | '.join(['---'] * len(headers))} |\n"
    
    for row in clean_table[1:]:
        markdown += f"| {' | '.join(row)} |\n"
    return markdown

def table_to_html(table: List[List[Any]]) -> str:
    if not table or not table[0]:
        return ""
    clean_table = []
    for row in table:
        clean_table.append([str(cell or "").strip().replace("\n", " ") for cell in row])
        
    headers = clean_table[0]
    if not any(headers):
        headers = [f"Column {i+1}" for i in range(len(headers))]
        
    html = "<table border='1' class='extracted-table'>\n  <thead>\n    <tr>\n"
    for h in headers:
        html += f"      <th>{h}</th>\n"
    html += "    </tr>\n  </thead>\n  <tbody>\n"
    
    for row in clean_table[1:]:
        html += "    <tr>\n"
        for cell in row:
            html += f"      <td>{cell}</td>\n"
        html += "    </tr>\n"
    html += "  </tbody>\n</table>"
    return html

def extract_pdf_split_router(file_path: str, gemini_api_key: Optional[str] = None) -> ExtractedContent:
    import fitz  # PyMuPDF
    import pdfplumber
    
    content = ExtractedContent()
    doc_pdf = fitz.open(file_path)
    total_pages = len(doc_pdf)
    content.stats.total_pages = total_pages
    
    table_ref_counter = 1
    figure_ref_counter = 1
    
    with pdfplumber.open(file_path) as pdf:
        for page_idx in range(total_pages):
            page_num = page_idx + 1
            fitz_page = doc_pdf[page_idx]
            plumber_page = pdf.pages[page_idx]
            
            # Pre-flight check: extract text length
            raw_text = fitz_page.get_text() or ""
            clean_text_len = len(raw_text.strip())
            
            # Fast-path if text length > 100
            if clean_text_len > 100:
                # Fast CPU Path
                # Extract tables using pdfplumber
                try:
                    tables = plumber_page.extract_tables()
                    if tables:
                        for tbl in tables:
                            md_tbl = table_to_markdown(tbl)
                            html_tbl = table_to_html(tbl)
                            if html_tbl:
                                ref_id = f"TABLE_REF_{table_ref_counter:03d}"
                                table_ref_counter += 1
                                
                                tbl_block = TableBlock(
                                    html_content=html_tbl,
                                    reference_id=ref_id,
                                    page_number=page_num,
                                    table_index=content.stats.table_count + 1,
                                    raw_markdown=md_tbl
                                )
                                content.table_chunks.append(tbl_block)
                                content.placeholders[ref_id] = html_tbl
                                content.stats.table_count += 1
                                
                                # Anchor placeholder into text
                                raw_text += f"\n\n[[{ref_id}]]\n\n"
                except Exception as tbl_err:
                    print(f"[FastPath] Table extraction error on p.{page_num}: {tbl_err}")
                
                # Check for standalone figures / images
                try:
                    image_list = fitz_page.get_images(full=True)
                    if image_list:
                        content.stats.image_count += len(image_list)
                except Exception:
                    pass
                    
                content.narrative_chunks.append(NarrativeBlock(
                    text=raw_text,
                    page_number=page_num
                ))
            else:
                # Visual Escalation Path (Scanned or visually dense page)
                content.stats.visual_pages += 1
                page_text = raw_text
                
                # If Gemini API Key is available, perform full page OCR & visual layout analysis
                if gemini_api_key:
                    try:
                        import base64
                        import requests
                        
                        pix = fitz_page.get_pixmap(dpi=200)
                        img_bytes = pix.tobytes("png")
                        img_b64 = base64.b64encode(img_bytes).decode('utf-8')
                        
                        gemini_url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={gemini_api_key}"
                        payload = {
                            "contents": [{
                                "parts": [
                                    {"text": "Perform high-precision OCR on this document page. Reconstruct all narrative text, section headings, and tabular data into clean Markdown formatting. Describe any key visual figures or charts present."},
                                    {
                                        "inlineData": {
                                            "mimeType": "image/png",
                                            "data": img_b64
                                        }
                                    }
                                ]
                            }]
                        }
                        headers = {"Content-Type": "application/json"}
                        res = requests.post(gemini_url, json=payload, headers=headers, timeout=30)
                        if res.status_code == 200:
                            res_data = res.json()
                            ocr_result = res_data["candidates"][0]["content"]["parts"][0]["text"]
                            page_text = f"[Scanned Page {page_num} OCR & Visual Parsing (Gemini 2.5 Flash)]\n" + ocr_result
                    except Exception as gemini_err:
                        print(f"[VisualPath] Gemini OCR failed on page {page_num}: {gemini_err}")
                
                content.narrative_chunks.append(NarrativeBlock(
                    text=page_text,
                    page_number=page_num
                ))
                
            content.stats.pages_processed += 1
            
    doc_pdf.close()
    return content

def extract_docx(file_path: str) -> ExtractedContent:
    import docx
    content = ExtractedContent()
    doc = docx.Document(file_path)
    
    full_text = []
    current_heading = None
    table_ref_counter = 1
    
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            current_heading = para.text
            full_text.append(f"\n# {para.text}\n")
        elif para.text.strip():
            full_text.append(para.text)
            
    for tbl_idx, table in enumerate(doc.tables):
        raw_table = []
        for row in table.rows:
            raw_table.append([cell.text.strip() for cell in row.cells])
            
        html_tbl = table_to_html(raw_table)
        md_tbl = table_to_markdown(raw_table)
        
        if html_tbl:
            ref_id = f"TABLE_REF_{table_ref_counter:03d}"
            table_ref_counter += 1
            
            content.table_chunks.append(TableBlock(
                html_content=html_tbl,
                reference_id=ref_id,
                page_number=1,
                table_index=tbl_idx + 1,
                raw_markdown=md_tbl
            ))
            content.placeholders[ref_id] = html_tbl
            content.stats.table_count += 1
            full_text.append(f"\n[[{ref_id}]]\n")
            
    text_str = "\n".join(full_text)
    content.narrative_chunks.append(NarrativeBlock(
        text=text_str,
        page_number=1,
        section_heading=current_heading
    ))
    content.stats.total_pages = 1
    content.stats.pages_processed = 1
    return content

def extract_excel_csv(file_path: str) -> ExtractedContent:
    import pandas as pd
    content = ExtractedContent()
    
    table_ref_counter = 1
    full_text = []
    
    if file_path.lower().endswith('.csv'):
        df_dict = {"Sheet1": pd.read_csv(file_path)}
    else:
        df_dict = pd.read_excel(file_path, sheet_name=None)
        
    for sheet_name, df in df_dict.items():
        # Clean NaNs
        df = df.fillna("")
        raw_table = [list(df.columns)] + df.values.tolist()
        
        html_tbl = table_to_html(raw_table)
        md_tbl = table_to_markdown(raw_table)
        
        ref_id = f"TABLE_REF_{table_ref_counter:03d}"
        table_ref_counter += 1
        
        content.table_chunks.append(TableBlock(
            html_content=html_tbl,
            reference_id=ref_id,
            page_number=1,
            table_index=content.stats.table_count + 1,
            raw_markdown=md_tbl
        ))
        content.placeholders[ref_id] = html_tbl
        content.stats.table_count += 1
        
        full_text.append(f"## Sheet: {sheet_name}\n[[{ref_id}]]\n")
        
    content.narrative_chunks.append(NarrativeBlock(
        text="\n".join(full_text),
        page_number=1
    ))
    content.stats.total_pages = 1
    content.stats.pages_processed = 1
    return content

def extract_pptx(file_path: str) -> ExtractedContent:
    import pptx
    content = ExtractedContent()
    prs = pptx.Presentation(file_path)
    
    total_slides = len(prs.slides)
    content.stats.total_pages = total_slides
    
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                slide_text.append(shape.text_frame.text)
                
        text_str = f"--- Slide {slide_num} ---\n" + "\n".join(slide_text)
        content.narrative_chunks.append(NarrativeBlock(
            text=text_str,
            page_number=slide_num
        ))
        content.stats.pages_processed += 1
        
    return content

def extract_html_md(file_path: str) -> ExtractedContent:
    content = ExtractedContent()
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        raw_content = f.read()
        
    if file_path.lower().endswith('.html') or file_path.lower().endswith('.htm'):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw_content, 'html.parser')
        text_str = soup.get_text(separator="\n")
    else:
        text_str = raw_content
        
    content.narrative_chunks.append(NarrativeBlock(
        text=text_str,
        page_number=1
    ))
    content.stats.total_pages = 1
    content.stats.pages_processed = 1
    return content
